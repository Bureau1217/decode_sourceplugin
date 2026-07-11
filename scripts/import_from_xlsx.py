#!/usr/bin/env python3
"""
Convert the DECODE Project 3 workbook's "Source Database - To fill" sheet
into public/data/sources.json.

Usage:
    python3 scripts/import_from_xlsx.py path/to/DECODE_Project_3_Source_Reliability_Plugin.xlsx

Requires: openpyxl (pip install openpyxl)

What it does:
  - Reads the sheet's header row (row 5) and data (from row 6).
  - Skips rows with no source_domain (the plugin can't match those to a
    search result) and prints their source_id so they can be fixed upstream.
  - Cleans source_domain: strips http(s)://, leading www., anything from
    the first "/" onward, and lowercases.
  - Parses authority_sources by splitting on whitespace and keeping tokens
    that look like URLs (a few rows have two URLs separated only by a
    space, with no comma/newline).
  - Maps topic_area -> topic, trust_score_100 -> trust_score, keeps the
    rest close to the sheet's own column names.
  - Writes public/data/sources.json wrapped in a Kirby-style API envelope.
"""

import sys
import re
import json
import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

SHEET_NAME = "Source Database - To fill"
HEADER_ROW = 5
FIRST_DATA_ROW = 6
OUT_PATH = Path(__file__).resolve().parent.parent / "public" / "data" / "sources.json"


def clean_domain(raw):
    if not raw:
        return None
    s = str(raw).strip()
    s = re.sub(r"^https?://", "", s, flags=re.I)
    s = re.sub(r"^www\.", "", s, flags=re.I)
    s = s.split("/")[0]
    return s.lower().strip()


def parse_urls(raw):
    if not raw:
        return []
    tokens = re.split(r"\s+", str(raw).strip())
    urls = []
    for t in tokens:
        t = t.strip().rstrip(".,;")
        if t.lower().startswith("http"):
            urls.append(t)
    return urls


def iso(d):
    if isinstance(d, datetime.datetime):
        return d.date().isoformat()
    if isinstance(d, datetime.date):
        return d.isoformat()
    return d


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[HEADER_ROW]]
    header = [h for h in header if h]

    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c + 1).value for c in range(len(header))]
        if any(v is not None for v in vals):
            rows.append(dict(zip(header, vals)))

    skipped = []
    out_rows = []
    for row in rows:
        domain = clean_domain(row.get("source_domain"))
        if not domain:
            skipped.append(row.get("source_id"))
            continue
        out_rows.append(
            {
                "source_id": row.get("source_id"),
                "source_name": row.get("source_name"),
                "source_domain": domain,
                "source_type": row.get("source_type"),
                "country_or_region": row.get("country_or_region"),
                "language": row.get("language"),
                "topic": row.get("topic_area"),
                "trust_score": row.get("trust_score_100"),
                "trust_score_raw": row.get("trust_score"),
                "score_scale": row.get("score_scale"),
                "trust_level": row.get("trust_level"),
                "trust_score_explanation": row.get("trust_score_explanation")
                or "No explanation provided yet.",
                "editorial_standards": row.get("editorial_standards"),
                "transparency": row.get("transparency"),
                "correction_policy": row.get("correction_policy"),
                "ownership_funding": row.get("ownership_funding"),
                "bias_notes": row.get("bias_notes"),
                "authority_sources": parse_urls(row.get("authority_sources")),
                "last_checked_date": iso(row.get("last_checked_date")),
                "status": row.get("status"),
                "student_owner": row.get("student_owner"),
                "validated_count": 0,
                "invalidated_count": 0,
            }
        )

    envelope = {"code": 200, "status": "ok", "result": {"data": out_rows}}
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(envelope, indent=2, ensure_ascii=False))

    print(f"written rows: {len(out_rows)}")
    print(f"skipped (no domain): {skipped}")
    print(f"output: {OUT_PATH}")


if __name__ == "__main__":
    main()
